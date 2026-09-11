import sys
import subprocess
sys.stdout.reconfigure(encoding='utf-8')


def install(package):
    subprocess.check_call([sys.executable, "-m", "pip", "install", package])

try:
    import pptx
except ImportError:
    install('python-pptx')
    import pptx

try:
    import PyPDF2
except ImportError:
    install('PyPDF2')
    import PyPDF2

def read_pptx(filepath, out_f):
    try:
        prs = pptx.Presentation(filepath)
        for slide in prs.slides:
            for shape in slide.shapes:
                if hasattr(shape, "text"):
                    out_f.write(shape.text + "\n")
    except Exception as e:
        out_f.write(f"Error reading {filepath}: {e}\n")

def read_pdf(filepath, out_f):
    try:
        with open(filepath, "rb") as f:
            reader = PyPDF2.PdfReader(f)
            for i in range(len(reader.pages)):
                out_f.write(reader.pages[i].extract_text() + "\n")
    except Exception as e:
        out_f.write(f"Error reading {filepath}: {e}\n")

with open("extract_output_direct.txt", "w", encoding="utf-8") as out_f:
    out_f.write("=== demo_1.pptx ===\n")
    read_pptx(r"d:\major_project\demo_1.pptx", out_f)
    out_f.write("=== demo_2.pptx ===\n")
    read_pptx(r"d:\major_project\demo_2.pptx", out_f)
    out_f.write("=== Reach_paper.pdf ===\n")
    read_pdf(r"d:\major_project\Reach_paper.pdf", out_f)

